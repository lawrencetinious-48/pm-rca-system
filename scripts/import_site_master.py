"""Import Airtel site master data from Lawrence Airtel.xlsx into config/site_master.json."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(" 00:00:00"):
        return text[:-9]
    return text


def sheet_records(workbook, sheet_name, header_row):
    worksheet = workbook[sheet_name]
    headers = [clean(value) for value in next(worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    records = []
    for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        record = {
            headers[index]: clean(value)
            for index, value in enumerate(row)
            if index < len(headers) and headers[index]
        }
        if any(record.values()):
            records.append(record)
    return records


def build_site_master(source_path: Path):
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    sites = {}

    for record in sheet_records(workbook, "SITE PM", 2):
        site_id = record.get("SITE ID", "").upper()
        if not site_id:
            continue
        sites.setdefault(site_id, {})
        sites[site_id].update(
            {
                "site_id": site_id,
                "site_name": record.get("SITE NAME", ""),
                "area": record.get("AREA", ""),
                "region": record.get("REGION", ""),
                "contractor": record.get("C0NTRACTOR", ""),
                "site_type": record.get("SITE TYPE", ""),
                "tower_owner": record.get("Tower Owner", ""),
                "power_share": record.get("Power Share", ""),
                "indoor_outdoor": record.get("Indoor/Outdoor", ""),
                "site_shelter": record.get("Site Shelter", ""),
                "guard_at_site": record.get("GUARD AT SITE", ""),
                "grid_available": record.get("Grid available", ""),
                "grid_readings": record.get("Grid Readings", ""),
                "site_category": record.get("Site Category", ""),
                "hybrid_available": record.get("Hybrid available(IPMU/INALA/SOLAR)", ""),
                "last_maintenance_date": record.get("Actual Date of Maintenance", ""),
                "noc_ref": record.get("NOC REF", ""),
                "last_technician": record.get("Name of Technician", ""),
                "technician_mobile": record.get("Technician Mobile Number", ""),
            }
        )

    for record in sheet_records(workbook, "New Sites", 2):
        site_id = record.get("Site_Id", "").upper()
        if not site_id:
            continue
        site = sites.setdefault(site_id, {"site_id": site_id})
        site["site_name"] = site.get("site_name") or record.get("Location Name", "")
        site.update(
            {
                "latitude": record.get("RRRU Lat", ""),
                "longitude": record.get("RRRU Long", ""),
                "frequency": record.get("Freq", ""),
                "mother_site_id": record.get("Mother Site ID", ""),
                "mother_site_name": record.get("Mother site Name", ""),
                "assigned_tech": record.get("Tech Name", ""),
            }
        )

    for record in sheet_records(workbook, "RECTIFIER & BATTERY TYPE", 1):
        site_id = record.get("SITE ID", "").upper()
        if not site_id:
            continue
        site = sites.setdefault(site_id, {"site_id": site_id})
        site["site_name"] = site.get("site_name") or record.get("SITE NAME", "")
        site.update(
            {
                "rectifier_type": record.get("Rectifier Type", ""),
                "rectifier_module_capacity": record.get("Rectifier Module Capacity", ""),
                "rectifier_modules": record.get("No. of Rectifier Modules", ""),
                "faulty_rectifier_modules": record.get("No. of Faulty Rectifier Modules", ""),
                "battery_type": record.get("Battery Type", ""),
                "battery_capacity": record.get("Battery Capacity", ""),
                "battery_strings": record.get("No. of Battery Strings", ""),
                "battery_count": record.get("No. of Batteries", ""),
                "battery_backup_time": record.get("Approximate Battery Backup Time", ""),
            }
        )

    issues_by_site = {}
    for record in sheet_records(workbook, "ISSUE TRACKER", 1):
        site_id = record.get("SITE ID", "").upper()
        if not site_id:
            continue
        issues_by_site.setdefault(site_id, []).append(
            {
                "issue": record.get("Issues picked", ""),
                "status": record.get("Status of Issue", ""),
                "noc_ref": record.get("Site Ref from Noc", ""),
                "responsibility": record.get("Responsibility", ""),
                "captured_date": record.get("Date Issue Captured", ""),
                "quantity": record.get("Quantity", ""),
                "eta": record.get("ETA", ""),
            }
        )

    for site_id, issues in issues_by_site.items():
        site = sites.setdefault(site_id, {"site_id": site_id})
        site["issues"] = issues
        site["open_issue_count"] = sum(
            1 for issue in issues if issue.get("status", "").lower() not in {"closed", "done", "resolved"}
        )

    return sorted(sites.values(), key=lambda item: item.get("site_id", ""))


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python scripts/import_site_master.py <Lawrence Airtel.xlsx> [output_json]")
    source_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("config/site_master.json")
    sites = build_site_master(source_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps({"source": source_path.name, "site_count": len(sites), "sites": sites}, indent=2),
        encoding="utf-8",
    )
    print(f"Wrote {len(sites)} sites to {output_path}")


if __name__ == "__main__":
    main()
